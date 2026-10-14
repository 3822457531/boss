"""L1 事实层 faq.md 读写与导入预览。"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from admin.paths import BACKUP_DIR, FAQ_PATH
from admin.audit import audit

_HEADING = re.compile(r"^##\s+(.+)\s*$")


def parse_faq_markdown(text: str) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    question: str | None = None
    body: list[str] = []
    for line in text.splitlines():
        m = _HEADING.match(line)
        if m:
            if question is not None:
                items.append({"question": question, "answer": "\n".join(body).strip()})
            question = m.group(1).strip()
            body = []
        else:
            if question is not None:
                body.append(line)
    if question is not None:
        items.append({"question": question, "answer": "\n".join(body).strip()})
    return [it for it in items if it["question"]]


def faq_to_markdown(items: list[dict[str, str]]) -> str:
    parts: list[str] = []
    for it in items:
        q = (it.get("question") or "").strip()
        a = (it.get("answer") or "").strip()
        if not q:
            continue
        parts.append(f"## {q}\n{a}\n")
    return "\n".join(parts).rstrip() + ("\n" if parts else "")


def load_faq(path: Path | None = None) -> list[dict[str, str]]:
    target = resolve_faq_path(str(path) if path else None)
    if not target.exists():
        return []
    return parse_faq_markdown(target.read_text(encoding="utf-8"))


def resolve_faq_path(rel: str | None = None) -> Path:
    from admin.paths import ROOT

    if rel is None or str(rel).strip() == "":
        return FAQ_PATH
    p = Path(rel)
    if p.is_absolute():
        return p
    return ROOT / p


def save_faq(
    items: list[dict[str, str]],
    *,
    actor: str = "local",
    path: Path | str | None = None,
) -> int:
    target = resolve_faq_path(str(path) if path else None)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        stamp = __import__("datetime").datetime.now().strftime("%Y%m%d-%H%M%S")
        bak = BACKUP_DIR / f"faq-{stamp}.md"
        bak.write_text(target.read_text(encoding="utf-8"), encoding="utf-8")
    text = faq_to_markdown(items)
    tmp = target.with_suffix(".md.tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(target)
    audit("faq_save", {"count": len(items), "path": str(target)}, actor=actor)
    return len(items)


def preview_from_upload(filename: str, raw: bytes) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(".md") or name.endswith(".markdown") or name.endswith(".txt"):
        return parse_faq_markdown(raw.decode("utf-8", errors="replace"))
    if name.endswith(".csv"):
        return _from_csv(raw.decode("utf-8-sig", errors="replace"))
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _from_xlsx(raw)
    if name.endswith(".docx"):
        return _from_docx(raw)
    # 默认当 markdown
    return parse_faq_markdown(raw.decode("utf-8", errors="replace"))


def preview_from_paste(text: str) -> list[dict[str, str]]:
    return parse_faq_markdown(text)


def _from_csv(text: str) -> list[dict[str, str]]:
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    start = 0
    if rows and any(h in "".join(rows[0]).lower() for h in ("问题", "question", "q")):
        start = 1
    items: list[dict[str, str]] = []
    for row in rows[start:]:
        if len(row) < 2:
            continue
        q, a = row[0].strip(), row[1].strip()
        if q:
            items.append({"question": q, "answer": a})
    return items


def _from_xlsx(raw: bytes) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("需要 openpyxl：pip install openpyxl") from exc
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    items: list[dict[str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if not row or row[0] is None:
            continue
        if i == 0 and str(row[0]).lower() in {"问题", "question", "q"}:
            continue
        q = str(row[0]).strip()
        a = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if q:
            items.append({"question": q, "answer": a})
    return items


def _from_docx(raw: bytes) -> list[dict[str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("需要 python-docx：pip install python-docx") from exc
    doc = Document(io.BytesIO(raw))
    # 用段落拼成 markdown：Heading 当问题
    lines: list[str] = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style else ""
        text = (p.text or "").strip()
        if not text:
            continue
        if "Heading" in style or text.startswith("##"):
            lines.append(f"## {text.lstrip('#').strip()}")
        else:
            lines.append(text)
    return parse_faq_markdown("\n".join(lines))
