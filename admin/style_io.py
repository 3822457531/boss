"""L2 风格层：列表、导入过滤、版本备份与回滚。"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from admin.audit import audit
from admin.config_io import read_policy
from admin.paths import BACKUP_DIR, RAW_HISTORY_PATH, STYLE_PAIRS_PATH
from src.knowledge import QAPair, dedupe, extract_pairs, save_pairs

_SALARY_OR_PROMISE = re.compile(
    r"(\d+\s*[kK万]|保证入职|包过|一定能转正|不用面试|保证过试用)"
)
_MAX_VERSIONS = 5


def load_style_pairs() -> list[dict[str, Any]]:
    if not STYLE_PAIRS_PATH.exists():
        return []
    rows: list[dict[str, Any]] = []
    for i, line in enumerate(STYLE_PAIRS_PATH.read_text(encoding="utf-8").splitlines()):
        line = line.strip()
        if not line.startswith("{"):
            continue
        try:
            obj = json.loads(line)
        except json.JSONDecodeError:
            continue
        obj["_id"] = i
        rows.append(obj)
    return rows


def style_stats() -> dict[str, Any]:
    rows = load_style_pairs()
    by_intent: dict[str, int] = {}
    qualities: list[float] = []
    for r in rows:
        intent = str(r.get("intent") or "other")
        by_intent[intent] = by_intent.get(intent, 0) + 1
        try:
            qualities.append(float(r.get("quality") or 0))
        except (TypeError, ValueError):
            pass
    mtime = None
    if STYLE_PAIRS_PATH.exists():
        mtime = datetime.fromtimestamp(STYLE_PAIRS_PATH.stat().st_mtime).isoformat(timespec="seconds")
    return {
        "total": len(rows),
        "by_intent": by_intent,
        "quality_avg": round(sum(qualities) / len(qualities), 2) if qualities else None,
        "quality_lt_1": sum(1 for q in qualities if q < 1),
        "quality_1_to_2": sum(1 for q in qualities if 1 <= q < 2),
        "quality_ge_2": sum(1 for q in qualities if q >= 2),
        "last_modified": mtime,
    }


def delete_style_ids(ids: list[int]) -> int:
    rows = load_style_pairs()
    keep = [r for r in rows if r.get("_id") not in set(ids)]
    _write_pairs_dicts(keep)
    audit("style_delete", {"deleted": len(rows) - len(keep)})
    return len(keep)


def parse_raw_history(
    *,
    recruiter_roles: set[str] | None = None,
    candidate_roles: set[str] | None = None,
    path: Path | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    """从 raw_history.jsonl 挖配对；返回 (pairs, sample_raw_lines)。"""
    target = path or RAW_HISTORY_PATH
    if not target.exists():
        raise FileNotFoundError(f"找不到 {target}")
    recruiter_roles = {x.lower() for x in (recruiter_roles or {"recruiter", "you", "myself", "hr"})}
    candidate_roles = {x.lower() for x in (candidate_roles or {"candidate", "friend", "geek"})}
    sample_raw: list[str] = []
    pairs: list[QAPair] = []
    policy = read_policy()
    style = policy.get("style") or {}
    min_q = int(style.get("min_question_chars") or 4)
    min_r = int(style.get("min_reply_chars") or 8)
    red = list(policy.get("redlines") or [])

    for line in target.read_text(encoding="utf-8").splitlines():
        if not line.strip().startswith("{"):
            continue
        if len(sample_raw) < 3:
            sample_raw.append(line[:1200])
        try:
            conv = json.loads(line)
        except json.JSONDecodeError:
            continue
        messages = []
        for m in conv.get("messages") or []:
            role_raw = str(m.get("role") or m.get("from") or "").strip().lower()
            if role_raw in recruiter_roles:
                role = "recruiter"
            elif role_raw in candidate_roles:
                role = "candidate"
            elif role_raw in {"system", "unknown"}:
                role = "system"
            else:
                continue
            text = str(m.get("text") or m.get("content") or "").strip()
            if text:
                messages.append({"role": role, "text": text})
        pairs.extend(
            extract_pairs(
                messages,
                source=str(conv.get("candidate") or ""),
                date=str(conv.get("date") or ""),
                min_question_chars=1,
                min_reply_chars=1,
                red_lines=[],
            )
        )
    # 过滤在 filter_pairs 里做，这里先挖全量
    dicts = [_pair_dict(p) for p in pairs]
    return dicts, sample_raw


def parse_table_pairs(raw: bytes, filename: str, q_col: int, a_col: int) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    rows_data: list[list[str]] = []
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows_data = [list(r) for r in csv.reader(io.StringIO(text))]
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            rows_data.append([("" if c is None else str(c)) for c in row])
    else:
        raise ValueError("表格仅支持 CSV / Excel")
    out: list[dict[str, Any]] = []
    for i, row in enumerate(rows_data):
        if i == 0 and _looks_header(row):
            continue
        if max(q_col, a_col) >= len(row):
            continue
        q, a = row[q_col].strip(), row[a_col].strip()
        if q and a:
            out.append(
                {
                    "question": q,
                    "reply": a,
                    "source": "upload",
                    "date": "",
                    "intent": "other",
                    "quality": 2.0,
                    "enabled": True,
                }
            )
    return out


def parse_paste_dialogue(text: str, recruiter_tag: str, candidate_tag: str) -> list[dict[str, Any]]:
    """简单粘贴：每行以标签开头，如 [HR] / [候选人]。"""
    messages: list[dict[str, str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        low = line.lower()
        if line.startswith(candidate_tag) or low.startswith(candidate_tag.lower()):
            messages.append({"role": "candidate", "text": line[len(candidate_tag) :].lstrip(" :：")})
        elif line.startswith(recruiter_tag) or low.startswith(recruiter_tag.lower()):
            messages.append({"role": "recruiter", "text": line[len(recruiter_tag) :].lstrip(" :：")})
    pairs = extract_pairs(messages, source="paste", min_question_chars=1, min_reply_chars=1, red_lines=[])
    return [_pair_dict(p) for p in pairs]


def filter_pairs(
    pairs: list[dict[str, Any]],
    *,
    min_question_chars: int = 4,
    min_reply_chars: int = 8,
    drop_salary_or_promise: bool = True,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for p in pairs:
        q = str(p.get("question") or "").strip()
        a = str(p.get("reply") or "").strip()
        if len(q) < min_question_chars or len(a) < min_reply_chars:
            continue
        if drop_salary_or_promise and _SALARY_OR_PROMISE.search(a):
            continue
        out.append(p)
    return out


def activate_style_pairs(pairs: list[dict[str, Any]], *, actor: str = "local") -> dict[str, Any]:
    """备份当前库并写入新数据。"""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    version_path = None
    if STYLE_PAIRS_PATH.exists():
        version_path = BACKUP_DIR / f"style_pairs-{stamp}.jsonl"
        version_path.write_text(STYLE_PAIRS_PATH.read_text(encoding="utf-8"), encoding="utf-8")
        _trim_versions()
    objects = [
        QAPair(
            question=str(p.get("question") or ""),
            reply=str(p.get("reply") or ""),
            source=str(p.get("source") or ""),
            date=str(p.get("date") or ""),
            intent=str(p.get("intent") or "other"),
            quality=float(p.get("quality") or 0),
        )
        for p in pairs
        if p.get("enabled", True) and str(p.get("question") or "").strip()
    ]
    objects = dedupe(objects)
    count = save_pairs(objects, STYLE_PAIRS_PATH)
    meta = {
        "activated_at": datetime.now().isoformat(timespec="seconds"),
        "count": count,
        "backup": str(version_path) if version_path else None,
        "actor": actor,
    }
    (BACKUP_DIR / "style_last_activate.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    audit("style_activate", {"count": count, "backup": meta["backup"]}, actor=actor)
    return meta


def list_style_versions() -> list[dict[str, Any]]:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    versions: list[dict[str, Any]] = []
    for path in sorted(BACKUP_DIR.glob("style_pairs-*.jsonl"), reverse=True):
        n = sum(1 for line in path.read_text(encoding="utf-8").splitlines() if line.strip().startswith("{"))
        versions.append(
            {
                "file": path.name,
                "path": str(path),
                "mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
                "count": n,
            }
        )
    return versions[:_MAX_VERSIONS]


def rollback_style(filename: str) -> dict[str, Any]:
    path = BACKUP_DIR / Path(filename).name
    if not path.exists() or not path.name.startswith("style_pairs-"):
        raise FileNotFoundError("备份不存在")
    # 先备份当前
    if STYLE_PAIRS_PATH.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        (BACKUP_DIR / f"style_pairs-{stamp}-before-rollback.jsonl").write_text(
            STYLE_PAIRS_PATH.read_text(encoding="utf-8"), encoding="utf-8"
        )
    STYLE_PAIRS_PATH.write_text(path.read_text(encoding="utf-8"), encoding="utf-8")
    audit("style_rollback", {"from": path.name})
    return {"ok": True, "restored": path.name, "stats": style_stats()}


def _trim_versions() -> None:
    files = sorted(BACKUP_DIR.glob("style_pairs-*.jsonl"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in files[_MAX_VERSIONS:]:
        try:
            old.unlink()
        except OSError:
            pass


def _pair_dict(p: QAPair) -> dict[str, Any]:
    return {
        "question": p.question,
        "reply": p.reply,
        "source": p.source,
        "date": p.date,
        "intent": p.intent,
        "quality": p.quality,
        "enabled": True,
    }


def _write_pairs_dicts(rows: list[dict[str, Any]]) -> None:
    objects = [
        QAPair(
            question=str(r.get("question") or ""),
            reply=str(r.get("reply") or ""),
            source=str(r.get("source") or ""),
            date=str(r.get("date") or ""),
            intent=str(r.get("intent") or "other"),
            quality=float(r.get("quality") or 0),
        )
        for r in rows
    ]
    save_pairs(objects, STYLE_PAIRS_PATH)


def _looks_header(row: list[str]) -> bool:
    joined = "".join(row).lower()
    return any(x in joined for x in ("问题", "question", "回复", "reply", "answer"))
